import requests
import pdfplumber
import pandas as pd
import docx
import openpyxl
from bs4 import BeautifulSoup
from datetime import datetime
from base_scraper import BaseScraper
from pathlib import Path
import re
import time
from typing import Dict, List, Optional, Union
import logging
import json
import os

class DocumentScraper(BaseScraper):
    """Classe base per l'estrazione di dati da documenti in vari formati."""
    
    def __init__(self, config_path: str = "config/config.yaml"):
        super().__init__(config_path)
        self.raw_data_dir = Path(self.config['percorsi']['raw_data'])
        self.max_retries = self.config['parametri_scraping'].get('retry_attempts', 3)
        self.timeout = self.config['parametri_scraping'].get('timeout', 30)
        
    def _scarica_documento(self, url: str) -> Optional[str]:
        """Scarica un documento e restituisce il path locale."""
        nome_file = url.split('/')[-1]
        local_path = self.raw_data_dir / nome_file
        
        if local_path.exists():
            self.logger.info(f"Documento già presente: {local_path}")
            return str(local_path)
            
        for tentativo in range(self.max_retries):
            try:
                self.logger.info(f"Tentativo {tentativo + 1} di download da {url}")
                response = requests.get(url, timeout=self.timeout)
                response.raise_for_status()
                
                with open(local_path, 'wb') as f:
                    f.write(response.content)
                    
                self.logger.info(f"Documento scaricato con successo: {local_path}")
                return str(local_path)
                
            except requests.Timeout:
                self.logger.warning(f"Timeout al tentativo {tentativo + 1}")
                if tentativo < self.max_retries - 1:
                    time.sleep(2 ** tentativo)  # Exponential backoff
                continue
                
            except requests.RequestException as e:
                self.logger.error(f"Errore nella richiesta HTTP: {str(e)}")
                if tentativo < self.max_retries - 1:
                    time.sleep(2 ** tentativo)
                continue
                
            except Exception as e:
                self.logger.error(f"Errore imprevisto: {str(e)}")
                return None
                
        self.logger.error(f"Impossibile scaricare il documento dopo {self.max_retries} tentativi")
        return None

    def _estrai_testo_da_pdf(self, file_path: str) -> str:
        """Estrae il testo da un file PDF."""
        testo_completo = []
        try:
            with pdfplumber.open(file_path) as pdf:
                for pagina in pdf.pages:
                    testo = pagina.extract_text()
                    if testo:
                        testo_completo.append(testo)
            return "\n".join(testo_completo)
        except Exception as e:
            self.logger.error(f"Errore nell'estrazione del testo dal PDF: {str(e)}")
            return ""

    def _estrai_testo_da_docx(self, file_path: str) -> str:
        """Estrae il testo da un file DOCX."""
        try:
            doc = docx.Document(file_path)
            return "\n".join([paragrafo.text for paragrafo in doc.paragraphs])
        except Exception as e:
            self.logger.error(f"Errore nell'estrazione del testo dal DOCX: {str(e)}")
            return ""

    def _estrai_testo_da_xlsx(self, file_path: str) -> str:
        """Estrae il testo da un file XLSX."""
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            testo = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for row in ws.rows:
                    testo.append(" ".join(str(cell.value) for cell in row if cell.value))
            return "\n".join(testo)
        except Exception as e:
            self.logger.error(f"Errore nell'estrazione del testo dall'XLSX: {str(e)}")
            return ""

    def _estrai_testo_da_documento(self, file_path: str) -> str:
        """Estrae il testo da un documento in base alla sua estensione."""
        estensione = os.path.splitext(file_path)[1].lower()
        
        if estensione == '.pdf':
            return self._estrai_testo_da_pdf(file_path)
        elif estensione == '.docx':
            return self._estrai_testo_da_docx(file_path)
        elif estensione in ['.xlsx', '.xls']:
            return self._estrai_testo_da_xlsx(file_path)
        else:
            self.logger.error(f"Formato documento non supportato: {estensione}")
            return ""

    def _pulisci_dati_missione(self, missione: dict) -> dict:
        """Pulisce e standardizza i dati estratti."""
        # Rimuovi spazi extra
        for key, value in missione.items():
            if isinstance(value, str):
                missione[key] = value.strip()
        
        # Converti date
        for campo in ['data_inizio', 'data_fine']:
            if missione.get(campo):
                try:
                    data = datetime.strptime(missione[campo], '%d/%m/%Y')
                    missione[campo] = data.strftime('%Y-%m-%d')
                except ValueError:
                    missione[campo] = ""
        
        # Converti numeri
        if missione.get('personale_totale'):
            try:
                missione['personale_totale'] = int(re.sub(r'[^\d]', '', str(missione['personale_totale'])))
            except ValueError:
                missione['personale_totale'] = 0
                
        if missione.get('costo_totale'):
            try:
                # Rimuovi punti e sostituisci virgole con punti
                costo = str(missione['costo_totale']).replace('.', '').replace(',', '.')
                missione['costo_totale'] = float(costo)
            except ValueError:
                missione['costo_totale'] = 0.0
        
        return missione

    def _estrai_dati_da_testo(self, testo: str, patterns: Dict[str, str]) -> Dict:
        """Estrae i dati dal testo usando i pattern regex forniti."""
        dati = {}
        for campo, pattern in patterns.items():
            match = re.search(pattern, testo, re.IGNORECASE | re.MULTILINE)
            dati[campo] = match.group(1).strip() if match else ""
        return dati

    def estrai_dati(self) -> pd.DataFrame:
        """Metodo da implementare nelle classi figlie."""
        raise NotImplementedError("Le classi figlie devono implementare questo metodo") 